"""Rule tests derived directly from the spec sections 3-12.  Run: python -m unittest discover -s tests -v"""
import os
import sys
import tempfile
import unittest
from datetime import datetime, date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from agent import config as C, dates, excel, pipeline, rules  # noqa: E402
from agent.location import classify  # noqa: E402
from agent.models import RawJob, Job, SourceStatus  # noqa: E402

NOW = datetime(2026, 9, 24, 8, 0, tzinfo=dates.TZ)
DA_DESC = ("We are looking for a Data Analyst to analyze data, build dashboards and reports, define KPIs and "
           "provide insights to stakeholders. Requirements: 2+ years of experience with SQL and Tableau. "
           "You will work closely with Data Engineers and Product.")
NO_KW_DESC = ("We are looking for an analyst to analyze business performance, build reports and dashboards, "
              "define KPIs and provide insights and recommendations to management using Excel.")
DE_DESC = ("Build and maintain scalable data pipelines with Airflow, Spark and Kafka. Develop ETL processes, "
           "data infrastructure and data lake on AWS, CI/CD, Docker, Kubernetes. Strong Python and SQL, Scala.")


def raw(title="Data Analyst", desc=DA_DESC, loc="Tel Aviv, Israel", posted="2026-09-20", updated="", ctx=False,
        url="https://example.com/job/1"):
    return RawJob(title=title, company="Acme", source="Test", url=url, description=desc, location=loc,
                  posted_raw=posted, updated_raw=updated, israeli_context=ctx)


class Stage1Titles(unittest.TestCase):
    def test_listed_and_variants(self):
        for t in ["Data Analyst", "DATA ANALYST", "Data-Analyst", "Senior Data & Analytics Analyst",
                  "Data Analyst / BI Analyst", "Data Analyst/ית", "Data Analyst - אנליסט/ית נתונים",
                  "BI Analyst/ית", "Sr. Product Analyst", "אנליסטית נתונים", "דאטה אנליסטית", "אנליסט/ית BI",
                  "אנליסטית פיננסית", "UA Analyst", "Credit Risk Analyst", "BI & Data Analyst"]:
            self.assertIn(rules.title_class(t), ("listed", "variant"), t)

    def test_excluded_titles(self):
        for t in ["Data Engineer", "Senior Data Scientist", "BI Developer", "Analytics Engineer", "DBA",
                  "ETL Developer", "Machine Learning Engineer", "Data Architect", "Full Stack Engineer"]:
            self.assertEqual(rules.title_class(t), "excluded", t)
            self.assertFalse(rules.is_potential(t), t)

    def test_non_data_analyst_titles_found_in_live_runs(self):
        for t in ["SOC Analyst", "Malware Analyst", "Security Analyst", "System Analyst", "מנתח/ת מערכות",
                  "Data Science & AI Lead", "Data Analytics Architect"]:
            self.assertFalse(rules.is_potential(t), t)
        for t in ["Senior Cyber Data Analyst", "Security Data Analyst"]:
            self.assertTrue(rules.is_potential(t), t)

    def test_precision_cases_from_full_run(self):
        for t in ["Data Engineering Team Lead", "BI & Data Engineering Team Leader", "Data & AI Product Manager",
                  "Director of Data", "Senior Marketing Data & Growth Manager", "AI Data Lead",
                  "Security Researcher Team Lead", "Senior Cyber Security Researcher", "אנליסט /ית NOC",
                  "ראש/ת צוות BI", "Senior Analytics Engineer – AI Analytics Platform"]:
            self.assertFalse(rules.is_potential(t), t)
        for t in ["מנתח-ת נתונים סניור", "מנתחת נתונים", "Senior Marketing Analytics", "מיישם BI", "Senior Economist"]:
            self.assertTrue(rules.is_potential(t), t)

    def test_precision_cases_round2(self):
        for t in ["Failure Analysis and Quality Control Team Leader", "Systems Analysis Team Lead",
                  "Head of Systems Analysis Team", "Senior Static Timing Analysis (STA) Engineer (IL)",
                  "Security Operations Center Analyst", "אנליסט.ית NOC",
                  "Analytics Predictive Analytics Solution Engineer"]:
            self.assertFalse(rules.is_potential(t), t)
        for t in ["אנליסט.ית שיווק", "אנליסט /ית תפעולית", "דרוש.ה אנליסט.ית נתונים"]:
            self.assertTrue(rules.is_potential(t), t)

    def test_title_abroad(self):
        job, why, _ = pipeline.evaluate(raw(title="Analyst, Marketing Analytics (Bangkok Based, relocation provided)",
                                            loc="Tel Aviv-Yafo, Israel"), NOW)
        self.assertIsNone(job)
        job, _, _ = pipeline.evaluate(raw(title="Data Analyst (Tel Aviv)"), NOW)
        self.assertIsNotNone(job)

    def test_combined_titles(self):
        for t in ["BI Developer / Data Analyst", "Data Engineer & Analyst"]:
            self.assertEqual(rules.title_class(t), "combined", t)

    def test_candidate_nonstandard(self):
        self.assertEqual(rules.title_class("Reporting & Insights Specialist"), "candidate")
        self.assertIsNone(rules.title_class("Account Executive"))
        self.assertFalse(rules.is_potential("Backend Developer"))


class Stage3Keywords(unittest.TestCase):
    def test_case_insensitive(self):
        for d in ["experience with sql", "Sql is a must", "SQL", "BigQuery", "power bi", "PowerBI", "Pandas",
                  "data warehouse", "DWH", "databases", "Qlik Sense", "Snowflake", "PostgreSQL", "T-SQL"]:
            self.assertTrue(rules.find_keywords(d), d)

    def test_no_false_hits(self):
        for d in ["NoSQL only", "snowflakes falling", "pythonic", "we love data", "Excel wizard"]:
            self.assertFalse(rules.find_keywords(d), d)

    def test_hebrew_equivalents(self):
        self.assertTrue(rules.find_keywords("ניסיון בעבודה עם מסדי נתונים"))
        self.assertFalse(rules.find_keywords("ניסיון בעבודה בנתונים ובאקסל"))

    def test_title_without_keyword_is_rejected(self):
        job, why, _ = pipeline.evaluate(raw(title="Senior Data Analyst", desc=NO_KW_DESC), NOW)
        self.assertIsNone(job)
        self.assertEqual(why, "אין מילת מפתח בתיאור")


class Stage4Essence(unittest.TestCase):
    def test_da_mentions_engineers_ok(self):
        job, why, _ = pipeline.evaluate(raw(), NOW)
        self.assertIsNotNone(job, why)

    def test_engineering_role_with_analyst_word_rejected(self):
        ok, _ = rules.essence("BI Developer / Data Analyst", DE_DESC)
        self.assertFalse(ok)

    def test_combined_title_analytics_ok(self):
        ok, _ = rules.essence("BI Developer / Data Analyst", DA_DESC + " Ad-hoc analysis, trends, A/B tests.")
        self.assertTrue(ok)

    def test_data_engineer_excluded_even_with_keywords(self):
        job, why, _ = pipeline.evaluate(raw(title="Data Engineer", desc=DE_DESC), NOW)
        self.assertIsNone(job)


class Stage5Location(unittest.TestCase):
    def check(self, loc, expected, **kw):
        self.assertEqual(classify(loc, **kw)[0], expected, loc)

    def test_israel(self):
        for loc in ["Tel Aviv-Yafo", "Tel Aviv District, Israel", "Herzliya", "רמת גן", "Petah Tikva, Israel",
                    "Israel", "Remote, Israel", "צפון תל אביב", "Ness Ziona", "Glilot", "רמת גן / גבעתיים"]:
            self.check(loc, "israel")

    def test_excluded(self):
        for loc in ["Rehovot", "רחובות", "Rishon LeZion", "ראשון לציון", "Jerusalem, Israel", "ירושלים",
                    "Haifa", "Yokneam Illit, Northern District, Israel", "Beer Sheva", "באר שבע", "Ashdod",
                    "צפון", "דרום", "Northern District, Israel", "Southern District"]:
            self.check(loc, "excluded")

    def test_multi_location_with_allowed_city_kept(self):
        self.check("Tel Aviv | Jerusalem", "israel")

    def test_foreign(self):
        for loc in ["New York, NY", "London, UK", "Remote - US", "Berlin, Germany", "Austin, TX"]:
            self.check(loc, "foreign")

    def test_unknown_and_context(self):
        self.assertEqual(classify("", "")[0], "unknown")
        self.assertEqual(classify("", "")[1], C.UNKNOWN)
        self.check("", "israel", israeli_context=True)
        self.check("Remote", "israel", description="Join our team in Tel Aviv")
        self.check("", "israel", url="https://www.company.co.il/careers/123")


class Stage6Dates(unittest.TestCase):
    def test_parse(self):
        cases = {"2026-09-20": date(2026, 9, 20), "2026-09-20T23:30:00Z": date(2026, 9, 21),
                 "20/09/2026": date(2026, 9, 20), "Today": date(2026, 9, 24), "Yesterday": date(2026, 9, 23),
                 "3 days ago": date(2026, 9, 21), "2 weeks ago": date(2026, 9, 10),
                 "Posted 10 days ago": date(2026, 9, 14), "לפני 7 שעות": date(2026, 9, 24),
                 "פורסם לפני 10 ימים": date(2026, 9, 14), "לפני שבועיים": date(2026, 9, 10),
                 "אתמול": date(2026, 9, 23), "1 ימים": date(2026, 9, 23), "30+ days ago": date(2026, 8, 25),
                 1790000000000: date(2026, 9, 21)}
        for raw_, exp in cases.items():
            self.assertEqual(dates.parse_date(raw_, NOW), exp, raw_)

    def test_update_date_wins(self):
        self.assertEqual(dates.relevant_date("2026-08-01", "2026-09-22", NOW), date(2026, 9, 22))

    def test_old_rejected_unknown_kept(self):
        job, why, _ = pipeline.evaluate(raw(posted="2026-08-01"), NOW)
        self.assertIsNone(job)
        job, why, _ = pipeline.evaluate(raw(posted="2026-08-01", updated="2026-09-20"), NOW)
        self.assertIsNotNone(job)
        job, why, _ = pipeline.evaluate(raw(posted=""), NOW)
        self.assertIsNotNone(job)
        self.assertEqual(job.date_display, C.UNKNOWN_DATE)
        self.assertIsNone(job.date_iso)
        job, _, _ = pipeline.evaluate(raw(posted="2026-09-04"), NOW)  # exactly 20 days -> kept
        self.assertIsNotNone(job)


class Cumulative(unittest.TestCase):
    def mk(self, url, d):
        j, _, _ = pipeline.evaluate(raw(url=url, posted=d or "", title="Data Analyst " + url[-1]), NOW)
        return j

    def test_new_dedupe_and_20_day_pruning(self):
        state = {"jobs": [], "last_sent_date": None}
        a = self.mk("https://x.com/job/a", "2026-09-20")
        b = self.mk("https://x.com/job/b", None)
        new, cum, removed = pipeline.update_cumulative(state, [a, b], [], NOW)
        self.assertEqual(len(new), 2)
        state["jobs"] = [j.to_dict() for j in cum]
        # next run: same job again (with tracking params) -> not new
        a2 = self.mk("https://x.com/job/a?utm_source=li", "2026-09-20")
        new, cum, _ = pipeline.update_cumulative(state, [a2], [], NOW)
        self.assertEqual(len(new), 0)
        self.assertEqual(len(cum), 2)
        # 30 days later: dated job removed, undated job kept for verification
        later = datetime(2026, 10, 24, 8, 0, tzinfo=dates.TZ)
        new, cum, removed = pipeline.update_cumulative(state, [], [], later)
        self.assertEqual(removed, 1)
        self.assertEqual([j.date_iso for j in cum], [None])
        # a later run finds a reliable old date for the undated job -> removed
        new, cum, removed = pipeline.update_cumulative({"jobs": [j.to_dict() for j in cum]}, [],
                                                       [(b.key, None, date(2026, 8, 1))], NOW)
        self.assertEqual(removed, 1)


class ExcelFile(unittest.TestCase):
    def test_four_sheets(self):
        j, _, _ = pipeline.evaluate(raw(), NOW)
        sts = [SourceStatus("https://a.com/x", "A", "m", count=1), SourceStatus("https://a.com/y", "A", "m"),
               SourceStatus("https://b.com", "B", "m", status="נכשל", error="HTTP 403")]
        with tempfile.TemporaryDirectory() as d:
            p = excel.build(os.path.join(d, "r.xlsx"), [j], [j], sts)
            self.assertTrue(excel.verify(p, 1))
            from openpyxl import load_workbook
            wb = load_workbook(p)
            self.assertEqual(wb.sheetnames, ["משרות חדשות", "רשימה מצטברת", "מצב המקורות", "מקורות ייחודיים"])
            self.assertEqual([c.value for c in wb["משרות חדשות"][1]], excel.JOB_HEADERS)
            self.assertEqual(wb["מצב המקורות"].max_row, 4)          # every scanned link is a row
            self.assertEqual([r[0].value for r in wb["מקורות ייחודיים"].iter_rows(min_row=2)], ["A"])  # dedup
            # empty "new jobs" still produces a valid file with headers only
            p2 = excel.build(os.path.join(d, "r2.xlsx"), [], [j], sts)
            self.assertTrue(excel.verify(p2, 0))


class UserLinks(unittest.TestCase):
    def test_parse_and_dedupe(self):
        with tempfile.TemporaryDirectory() as d:
            old = C.USER_SOURCES_DIR
            C.USER_SOURCES_DIR = d
            try:
                with open(os.path.join(d, "a.txt"), "w", encoding="utf-8") as f:
                    f.write("\n  https://a.com/careers  \n# comment\nhttps://a.com/careers\n\nb.co.il/jobs\n"
                            "https://jobs.lever.co/x?utm_source=1\nhttps://jobs.lever.co/x\n")
                links = pipeline.read_user_links()
            finally:
                C.USER_SOURCES_DIR = old
        self.assertEqual(links, ["https://a.com/careers", "https://b.co.il/jobs", "https://jobs.lever.co/x?utm_source=1"])

    def test_same_site_different_path_scanned_once(self):
        with tempfile.TemporaryDirectory() as d:
            old = C.USER_SOURCES_DIR
            C.USER_SOURCES_DIR = d
            try:
                with open(os.path.join(d, "a.txt"), "w", encoding="utf-8") as f:
                    f.write("https://www.max.co.il/jobs/lobby\n")
                with open(os.path.join(d, "b.txt"), "w", encoding="utf-8") as f:
                    f.write("1. https://max.co.il/careers\n"
                            "https://www.comeet.com/jobs/a/11.111\nhttps://www.comeet.com/jobs/b/22.222\n")
                links = pipeline.read_user_links()
            finally:
                C.USER_SOURCES_DIR = old
        self.assertEqual(links, ["https://www.max.co.il/jobs/lobby", "https://www.comeet.com/jobs/a/11.111",
                                 "https://www.comeet.com/jobs/b/22.222"])

    def test_email_body_has_summary_but_no_job_list(self):
        from agent import mailer
        j, _, _ = pipeline.evaluate(raw(title="Unique Data Analyst Title"), NOW)
        body = mailer.body_html("25/09/2026", [j], [j], [SourceStatus("https://a.com", "A", "m")], 0)
        self.assertIn("משרות חדשות:</b> 1", body)
        self.assertIn("קישורים שנסרקו:</b> 1", body)
        self.assertNotIn("Unique Data Analyst Title", body)
        self.assertNotIn("<table", body)


class Timing(unittest.TestCase):
    def test_send_time_is_10_israel_in_summer_and_winter(self):
        from agent.main import at_il
        summer = at_il(date(2026, 7, 1), C.SEND_HOUR)
        winter = at_il(date(2026, 12, 1), C.SEND_HOUR)
        self.assertEqual(summer.utcoffset().total_seconds(), 3 * 3600)
        self.assertEqual(winter.utcoffset().total_seconds(), 2 * 3600)
        self.assertEqual((summer.hour, winter.hour), (10, 10))


class CompanyTypeColumn(unittest.TestCase):
    """The agreed "סוג חברה" column: classification, row order, position and frozen columns."""

    def test_classify_by_name(self):
        from agent import company_type as CT
        cases = {"Similarweb": CT.HITECH, "Palo Alto Networks": CT.HITECH, "בנק לאומי": CT.BANKS,
                 "Bank Hapoalim בנק הפועלים": CT.BANKS, "הראל ביטוח ופיננסים": CT.FINANCE,
                 "Cal (Israel Credit Cards)": CT.FINANCE, "EY": CT.CONSULTING, "G-STAT": CT.CONSULTING,
                 "Novo Nordisk": CT.HEALTH, "קופת חולים לאומית": CT.HEALTH, "Keshet Media Group": CT.MEDIA,
                 "פלאפון": CT.MEDIA, "אסם": CT.RETAIL, "Kimberly-Clark": CT.RETAIL, "Thales": CT.INDUSTRY,
                 "Enlight Renewable Energy Ltd (ENLT)": CT.INDUSTRY, "קבוצת ארקיע": CT.OTHER,
                 "החברה המאוחדת לתיירות בעמ": CT.OTHER}
        for company, expected in cases.items():
            self.assertEqual(CT.classify(company)[0], expected, company)

    def test_hidden_and_agencies(self):
        from agent import company_type as CT
        for company in ["חסוי", "- חסוי -", "Confidential- Discreet", "לא צוין במקור",
                        "חברה בתחום הייטק / חומרה / תוכנה / סייבר", "Experis Academy", "לין ביכלר",
                        "יערה פיינר אבחון והשמה בע\"מ", "Recruitx", ""]:
            self.assertEqual(CT.classify(company)[0], CT.AGENCY, company)

    def test_other_signals_when_name_is_unknown(self):
        from agent import company_type as CT
        self.assertEqual(CT.classify("Zorblax", industry="Software Development")[0], CT.HITECH)
        self.assertEqual(CT.classify("Zorblax", industry="Banking")[0], CT.BANKS)
        self.assertEqual(CT.classify("Zorblax", industry="Staffing and Recruiting")[0], CT.AGENCY)
        self.assertEqual(CT.classify("Zorblax", industry="Hospitals and Health Care")[0], CT.HEALTH)
        self.assertEqual(CT.classify("Zorblax", industry="Hospitality")[0], CT.OTHER)
        self.assertEqual(CT.classify("Zorblax", listing_url="https://job-boards.greenhouse.io/zorblax")[0], CT.HITECH)
        self.assertEqual(CT.classify("Zorblax", description="A fast growing startup building a SaaS product")[0],
                         CT.HITECH)
        self.assertEqual(CT.classify("Zorblax", description="We sell furniture")[0], CT.OTHER)

    def test_one_type_per_company(self):
        from agent import company_type as CT
        a = raw(title="Data Analyst", url="https://x.com/1")
        b = raw(title="BI Analyst", url="https://x.com/2")
        a.company = b.company = "Zorblax"
        a.industry = "Banking"
        ja, _, _ = pipeline.evaluate(a, NOW)
        jb, _, _ = pipeline.evaluate(b, NOW)
        types = CT.company_map([(ja, a), (jb, b)], pipeline._company_key)
        self.assertEqual(types, {pipeline._company_key("Zorblax"): CT.BANKS})

    def test_rows_sorted_by_agreed_order_newest_first_inside(self):
        from agent import company_type as CT

        def job(company, ctype, d):
            j, _, _ = pipeline.evaluate(raw(title="Data Analyst " + company, posted=d,
                                            url=f"https://x.com/{company}"), NOW)
            j.company = company
            j.company_type = ctype
            return j
        jobs = [job("r", CT.RETAIL, "2026-09-23"), job("a", CT.AGENCY, "2026-09-23"), job("h1", CT.HITECH, "2026-09-10"),
                job("b", CT.BANKS, "2026-09-22"), job("h2", CT.HITECH, "2026-09-22"), job("o", CT.OTHER, "2026-09-23")]
        jobs.sort(key=lambda j: j.date_iso, reverse=True)
        out = CT.sort_jobs(jobs)
        self.assertEqual([j.company for j in out], ["h2", "h1", "b", "r", "o", "a"])

    def test_cumulative_gets_types_and_order(self):
        from agent import company_type as CT
        state = {"jobs": []}
        bank = raw(title="Data Analyst", url="https://x.com/bank")
        bank.company = "בנק לאומי"
        tech = raw(title="Data Analyst", url="https://x.com/tech", posted="2026-09-10")
        tech.company = "Similarweb"
        jobs = [pipeline.evaluate(r, NOW)[0] for r in (bank, tech)]
        types = CT.company_map(list(zip(jobs, (bank, tech))), pipeline._company_key)
        for j in jobs:
            j.company_type = types[pipeline._company_key(j.company)]
        new, cum, _ = pipeline.update_cumulative(state, jobs, [], NOW, types)
        self.assertEqual([j.company_type for j in new], [CT.HITECH, CT.BANKS])  # tech first despite older date
        # an old saved job without a type (saved before this column existed) is classified by name
        old = {k: v for k, v in cum[1].to_dict().items() if k != "company_type"}
        _, cum2, _ = pipeline.update_cumulative({"jobs": [old]}, [], [], NOW)
        self.assertEqual(cum2[0].company_type, CT.BANKS)

    def test_excel_column_after_role_and_frozen_panes(self):
        from agent import company_type as CT
        from openpyxl import load_workbook
        j, _, _ = pipeline.evaluate(raw(), NOW)
        j.company_type = CT.HITECH
        with tempfile.TemporaryDirectory() as d:
            p = excel.build(os.path.join(d, "r.xlsx"), [j], [j], [SourceStatus("https://a.com", "A", "m")])
            wb = load_workbook(p)
            for name in ("משרות חדשות", "רשימה מצטברת"):
                ws = wb[name]
                header = [c.value for c in ws[1]]
                self.assertEqual(header[:3], ["תפקיד", "סוג חברה", "חברה"])
                self.assertEqual(len(header), 9)
                self.assertEqual(ws["B2"].value, CT.HITECH)
                self.assertEqual(ws.freeze_panes, "C2")
                self.assertEqual(ws.cell(row=2, column=8).hyperlink.target, j.url)
            self.assertEqual(wb["מצב המקורות"].freeze_panes, "A2")  # source tabs unchanged


if __name__ == "__main__":
    unittest.main()
