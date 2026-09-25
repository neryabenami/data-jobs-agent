"""Section 11: the .xlsx report with exactly 4 sheets."""
import os

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

SHEETS = ["משרות חדשות", "רשימה מצטברת", "מצב המקורות", "מקורות ייחודיים"]
JOB_HEADERS = ["תפקיד", "חברה", "מקור", "מיקום", "ניסיון", "סיבה להתאמה", "קישור ישיר למשרה", "תאריך פרסום המשרה"]
JOB_WIDTHS = [38, 24, 18, 24, 22, 60, 55, 22]
SOURCE_HEADERS = ["מקור / קישור שנסרק", "אתר", "שיטת איסוף", "סטטוס", "תיאור שגיאה",
                  "מספר משרות רלוונטיות שחולצו"]
SOURCE_WIDTHS = [70, 26, 40, 12, 50, 16]
UNIQUE_HEADERS = ["שם האתר"]

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
LINK_FONT = Font(color="0563C1", underline="single")
STATUS_FILL = {"הצלחה": "E2EFDA", "חלקי": "FFF2CC", "נכשל": "F8CBAD", "לא נגיש": "F8CBAD"}
THIN = Border(bottom=Side(style="thin", color="D9D9D9"))


def _sheet(ws, headers, widths):
    ws.sheet_view.rightToLeft = True
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(row=1, column=i)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30


def _jobs(ws, jobs):
    _sheet(ws, JOB_HEADERS, JOB_WIDTHS)
    for j in jobs:
        ws.append([j.title, j.company, j.source, j.location, j.experience, j.reason, j.url, j.date_display])
        r = ws.max_row
        link = ws.cell(row=r, column=7)
        if j.url.startswith("http"):
            link.hyperlink, link.font = j.url, LINK_FONT
        for col in range(1, 9):
            cell = ws.cell(row=r, column=col)
            cell.alignment = Alignment(vertical="top", wrap_text=col in (1, 6))
            cell.border = THIN
    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:H{ws.max_row}"


def build(path, new_jobs, cumulative, statuses):
    wb = Workbook()
    ws_new = wb.active
    ws_new.title = SHEETS[0]
    _jobs(ws_new, new_jobs)
    _jobs(wb.create_sheet(SHEETS[1]), cumulative)

    ws_src = wb.create_sheet(SHEETS[2])
    _sheet(ws_src, SOURCE_HEADERS, SOURCE_WIDTHS)
    for s in statuses:
        ws_src.append([s.url, s.site, s.method, s.status, s.error, s.count])
        r = ws_src.max_row
        if s.url.startswith("http"):
            ws_src.cell(row=r, column=1).hyperlink = s.url
            ws_src.cell(row=r, column=1).font = LINK_FONT
        ws_src.cell(row=r, column=4).fill = PatternFill("solid", fgColor=STATUS_FILL.get(s.status, "FFFFFF"))
        ws_src.cell(row=r, column=5).alignment = Alignment(wrap_text=True, vertical="top")
    if ws_src.max_row > 1:
        ws_src.auto_filter.ref = f"A1:F{ws_src.max_row}"

    ws_u = wb.create_sheet(SHEETS[3])
    _sheet(ws_u, UNIQUE_HEADERS, [40])
    for name in unique_sites(statuses):
        ws_u.append([name])

    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)
    return path


def unique_sites(statuses):
    """Sites that actually took part in the scan (at least one link scanned successfully or partially)."""
    seen = {}
    for s in statuses:
        if s.site.startswith(("Bing", "גילוי")):
            continue  # discovery tools find sources; they are not job sources themselves
        if s.status in ("הצלחה", "חלקי") and s.site and not s.error.startswith("לא נסרק"):
            seen.setdefault(s.site.lower(), s.site)
    return sorted(seen.values(), key=str.lower)


def verify(path, expected_new):
    """Sanity check before sending (section 12.3): real xlsx, exactly the 4 sheets, row counts match."""
    wb = load_workbook(path, read_only=True)
    names = wb.sheetnames
    if names != SHEETS:
        raise ValueError(f"גליונות לא תקינים: {names}")
    rows_new = wb[SHEETS[0]].max_row - 1
    if rows_new != expected_new:
        raise ValueError(f"מספר משרות חדשות בקובץ ({rows_new}) לא תואם ({expected_new})")
    wb.close()
    return True
